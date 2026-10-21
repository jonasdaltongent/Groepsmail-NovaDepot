#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het werkdocument voor les 12 "Aan, CC, BCC en een bijlage" (NovaDepot, ORLO, Communicatie 2).

  CO2_Aan-CC-BCC-en-een-bijlage.docx   het werkdocument dat de leerling INLEVERT (een kopie per leerling)

Geen bronbestanden en geen zip: de les werkt alleen in het werkdocument en in de browser. Geen Gmail: de leerlingen
hebben er met hun schoolaccount geen toegang toe (Jonas, 09-10-2026). De e-mail komt in dit document, in een
venster dat eruitziet als "Nieuw bericht", met Aan, Cc, Bcc, Onderwerp en Bijlage. Er wordt niets verstuurd.
De bijlage is de eigen pdf uit les 04: de leerling typt in het vak Bijlage @ en een stuk van de naam, en kiest de
pdf (een chip). Wie ze niet vindt, typt de naam.

Gebruik:  python3 maak_werkdocumenten.py
Vereist:  python-docx, Pillow (huisstijl.py en go-dalton-gent.png staan ernaast)

LET OP bij aanpassen:
 - Alle namen, nummers en adressen zijn verzonnen; de adressen eindigen op .example (RFC 2606).
 - De leerling schrijft VOOR Lotte (naam en handtekening van Lotte): zo staat er geen eigen naam of gegeven in het
   werkdocument.
 - Rij 1 van deel 1 is de demo van dia 5, al ingevuld in het grijs.
 - De punten staan per deel (4 · 3 · 13) en samen in de controletabel (lesvoorbereiding § 22). Verander je ze, doe
   het op beide plaatsen.
 - Eén vak per ding dat nagekeken wordt (_afspraken/verbeteren.md): Aan, Cc, Bcc, Onderwerp, Bijlage en de zes
   vakken van de tekst. De verwachte antwoorden staan in lesvoorbereiding § 14 en § 22.
 - Het deel 'Nalezen' begint met het vak voor het codewoord (PAPERCLIP); de leerling krijgt het in stap 5 van de
   lespagina na de zelftest.
"""
import os

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import RGBColor

import huisstijl as hs   # de huisstijl van het werkdocument (uit _sjablonen/lespakket/werkdocument/)

HERE = os.path.dirname(os.path.abspath(__file__))
UIT = os.path.join(HERE, "CO2_Aan-CC-BCC-en-een-bijlage.docx")

# Het kader met de feiten (lesvoorbereiding § 4). De vijf adressen staan elk op een eigen regel: om te kopiëren.
JOBSTUDENTEN = [
    "lina.haddad@mail.example",
    "emre.yilmaz@mail.example",
    "noor.verbeke@mail.example",
    "kobe.desmet@mail.example",
    "ayla.janssens@mail.example",
]
TOM = "tom.peeters@novadepot.example"
FEITEN = [
    ("Van:", "Lotte Maes, personeelsdienst van NovaDepot, toestel 205"),
    ("Voor:", "de vijf nieuwe jobstudenten. Ze kennen elkaar niet. Hun adressen zijn privé:"),
] + JOBSTUDENTEN + [
    ("Ter informatie:", "Tom Peeters, magazijnverantwoordelijke (moet niets doen)"),
    TOM,
    ("Bijlage:", "de onthaalbrochure: de pdf die jij in les 04 maakte"),
    ("Eerste werkdag:", "maandag 21 december 2026, om 7.30 uur aan het onthaal; breng je identiteitskaart mee"),
    ("Daarna:", "om 8.00 uur infosessie in de kantine"),
    ("Vraag van Lotte:", "lees de brochure, en laat tegen vrijdag 11 december weten dat je komt"),
]

# Deel 1: situatie · wie · veld. Rij 1 = de demo, al ingevuld in het grijs.
RIJEN_VELDEN = [
    ("Tom vraagt Kim Wouters van De Melkkan om een levering te verplaatsen.", "Kim Wouters",
     "Aan — zij moet iets doen"),
    ("Tom wil dat Sara Claes (preventie) weet dat de koelcel onderhoud krijgt. Sara moet niets doen.",
     "Sara Claes", ""),
    ("Het onthaal stuurt een nieuwjaarswens naar 40 klanten. De klanten kennen elkaar niet.",
     "de 40 klanten", ""),
    ("Lotte vraagt Tom om het uurrooster van de jobstudenten na te kijken en terug te sturen.", "Tom", ""),
    ("Tom stuurt de planning van de week naar de 8 collega's van zijn ploeg. Ze werken samen en moeten de planning "
     "lezen.", "de 8 collega's", ""),
]

# Deel 2: de vragen onder de mail van Lotte
RIJEN_ZIET = [
    ("a. Lina opent de mail. Ziet ze het adres van Tom? Schrijf ja of nee.", ""),
    ("b. Ziet Lina het adres van Emre? Schrijf ja of nee.", ""),
    ("c. Waarom zet Lotte de jobstudenten in Bcc? Schrijf één zin. Kort mag.", ""),
]

# De criteria van lesvoorbereiding § 22 (samen 20): 4 · 3 · 13
CRITERIA = [
    ("Deel 1: bij elke situatie staat het juiste veld: Aan, Cc of Bcc.", 4),
    ("Deel 2: ik schreef juist wie welk adres ziet (a en b).", 2),
    ("Deel 2: ik schreef waarom de jobstudenten in Bcc staan (c).", 1),
    ("Deel 3: Aan is leeg, Tom staat in Cc en de vijf jobstudenten staan in Bcc.", 3),
    ("Deel 3: mijn onderwerp zegt waarover de mail gaat.", 1),
    ("Deel 3: de brochure staat in het vak Bijlage, en mijn tekst zegt dat ze in de bijlage zit.", 2),
    ("Deel 3: ik begin met *Beste jobstudenten,* en sluit af met een slotgroet zonder afkorting. Mijn inleiding zegt "
     "wie schrijft en waarom.", 2),
    ("Deel 3: mijn boodschap heeft de eerste werkdag, het uur, het onthaal, de identiteitskaart en de infosessie.", 2),
    ("Deel 3: mijn slot zegt tegen wanneer ze antwoorden. Ik schrijf overal *u*, zonder chattaal of emoji's.", 2),
    ("Deel 3: onder mijn mail staat een handtekening met naam, personeelsdienst, NovaDepot en toestel 205.", 1),
]


def leesvenster(doc):
    """De mail van Lotte zoals ze in deel 2 staat: een venster dat de leerling alleen leest (grijze rand, geen
    lichtblauwe vakken). Aan is leeg, Cc heeft Tom, Bcc de vijf adressen; het onderwerp en de tekst zijn grijs: de
    mail zelf schrijft de leerling in deel 3."""
    velden = [
        ("Van", ["lotte.maes@novadepot.example"], False),
        ("Aan", ["(leeg)"], True),
        ("Cc", [TOM], False),
        ("Bcc", JOBSTUDENTEN, False),
        ("Onderwerp", ["…"], True),
    ]
    t = doc.add_table(rows=1 + len(velden) + 1, cols=2)
    t.alignment = hs.WD_TABLE_ALIGNMENT.CENTER
    hs.randen(t, buiten="9AA0A6")
    hs.breedtes(t, [3.2, hs.BREEDTE - 3.2])
    kopcel = t.rows[0].cells[0].merge(t.rows[0].cells[1])
    hs.schaduw(kopcel, "F2F6FC")
    hs.celtekst(kopcel, "Nieuw bericht", vet=True, kleur=RGBColor(0x20, 0x21, 0x24))
    kopcel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    hs.hoogte(t.rows[0], 0.8)
    for i, (veld, regels, grijs) in enumerate(velden, start=1):
        links, rechts = t.rows[i].cells
        hs.celtekst(links, veld, kleur=hs.GRIJS)
        hs.celtekst(rechts, regels[0], kleur=hs.GRIJS if grijs else None, cursief=grijs)
        for regel in regels[1:]:
            p = rechts.add_paragraph()
            hs.zonder_ruimte(p)
            hs.stijl(p.add_run(regel), None, None)
        for c in (links, rechts):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            hs.cel_randen(c, "DADCE0", dikte=6, kanten=("bottom",))
        hs.hoogte(t.rows[i], 0.85)
    body = t.rows[len(velden) + 1].cells[0].merge(t.rows[len(velden) + 1].cells[1])
    hs.schaduw(body, "F1F3F4")
    hs.marges(body, boven=0.25, onder=0.25, links=0.3, rechts=0.3)
    hs.celtekst(body, "De tekst van de mail staat hier. Die schrijf jij zelf in deel 3.", kleur=hs.GRIJS, cursief=True)
    hs.hoogte(t.rows[len(velden) + 1], 1.6)
    hs.niet_splitsen(t)
    hs.ruimte(doc, 4)
    return t


def nieuwe_bladzijde(doc):
    """Een bladzijdeskeer vóór een deel met puntenvak: page_break_before in een tabelcel (huisstijl.deel met punten)
    negeert Google Documenten. Een kleine lege alinea die zelf op een nieuwe bladzijde begint, doet het wel."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    hs.zonder_ruimte(p)
    p.paragraph_format.line_spacing = hs.Pt(2)
    r = p.add_run("")
    r.font.size = hs.Pt(1)
    return p


def werkdocument():
    """In de huisstijl van 12-10-2026 (huisstijl.py): alleen wat de leerling invult, en de feiten die hij
    daarvoor nodig heeft. Hoe hij werkt, staat op de lespagina."""
    doc = hs.nieuw_document()
    doc.core_properties.title = "CO2 Aan, CC, BCC en een bijlage - werkdocument"
    doc.core_properties.author = "Toegepaste Informatica"
    hs.kopblok(doc, "AAN, CC, BCC EN EEN BIJLAGE", punten=20)

    hs.kader(doc, "Lotte mailt de onthaalbrochure naar de nieuwe jobstudenten", [
        "Lotte Maes (personeelsdienst) vraagt jou: mail de onthaalbrochure naar de vijf nieuwe jobstudenten. Zet Tom "
        "in kopie: hij is hun verantwoordelijke in het magazijn. Ik lees de mail na en verstuur ze zelf.",
    ] + FEITEN)

    hs.deel(doc, 1, "Aan, Cc of Bcc?", punten=4)
    hs.tabel(doc, ["Situatie", "Wie?", "Aan, Cc of Bcc?"], RIJEN_VELDEN, [8.6, 3.0, hs.BREEDTE - 11.6],
             invul=(2,), rijhoogte=1.5, voorbeeld_rij=1)

    nieuwe_bladzijde(doc)
    hs.deel(doc, 2, "Wie ziet welk adres?", punten=3)
    hs.tekst(doc, "Dit is de mail van Lotte aan de jobstudenten. Hij is nog niet verstuurd.", klein=True, cursief=True)
    leesvenster(doc)
    hs.tabel(doc, ["Vraag", "Antwoord"], RIJEN_ZIET, [10.6, hs.BREEDTE - 10.6], invul=(1,), rijhoogte=1.5)

    nieuwe_bladzijde(doc)
    hs.deel(doc, 3, "De mail van Lotte", punten=13)
    hs.tekst(doc, "Schrijf zoals Lotte het zou versturen: u-vorm, geen afkortingen, geen emoji.", klein=True,
             cursief=True)
    venster = hs.mailvenster(doc, [("Aanspreking", 1.0), ("Inleiding", 2.0), ("De boodschap", 4.2), ("Slot", 2.0),
                                   ("Slotgroet", 1.0), ("Handtekening", 2.0)],
                             velden=("Aan", "Cc", "Bcc", "Onderwerp", "Bijlage"))
    hs.hoogte(venster.rows[3], 2.3)   # Bcc: vijf adressen
    hs.hoogte(venster.rows[5], 1.2)   # Bijlage: een chip met een lange bestandsnaam

    hs.deel(doc, 4, "Nalezen", nieuwe_pagina=True)
    # eerst het codewoord van de zelftest (15-10-2026): de lespagina toont het pas als de drie vragen van stap 5
    # juist zijn; zo ziet Jonas wie de zelftest echt maakte (_afspraken/lespagina.md, De laatste stap)
    hs.vraag(doc, "Het codewoord van de zelftest (stap 5 op de lespagina):")
    hs.antwoordvak(doc, 0.9)
    hs.controletabel(doc, CRITERIA)

    hs.deel(doc, None, "Extra — niet verplicht")
    hs.vraag(doc, "Je pdf heet nog zoals in les 04. Welke naam geef je ze voor de jobstudenten? Schrijf de naam "
                  "volgens de afspraak van les 06 (Wat - Van wie).")
    hs.antwoordvak(doc, 1.6)

    doc.save(UIT)


if __name__ == "__main__":
    werkdocument()
    print("Klaar.")
    print("  werkdocument om in te leveren : " + os.path.basename(UIT))
